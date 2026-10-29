"""Демо-колода PyDungeon: 3 слайда, 16:9, визуальный язык DESIGN.md.

Запуск:  python demo/deck/build_deck.py
Рендер:  powershell -File demo/deck/render.ps1
"""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from notes import NOTES

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
ASSETS = HERE / "assets"
OUT = HERE / "pydungeon-deck.pptx"

# --- палитра DESIGN.md -------------------------------------------------------
NIGHT = RGBColor(0x09, 0x0B, 0x0A)
NIGHT2 = RGBColor(0x0F, 0x14, 0x11)
PAPER = RGBColor(0xD9, 0xCB, 0xB0)
PAPER2 = RGBColor(0xC8, 0xB7, 0x9A)
PAPER3 = RGBColor(0xE6, 0xDC, 0xC6)
INK = RGBColor(0x1C, 0x1A, 0x16)
INK_SOFT = RGBColor(0x48, 0x44, 0x3A)
RUST = RGBColor(0x8A, 0x4A, 0x32)
RUST_INK = RGBColor(0x71, 0x38, 0x23)
RUST_LIGHT = RGBColor(0xD3, 0x9A, 0x74)

SERIF = "Literata"
MONO = "IBM Plex Mono"

W, H = 13.333, 7.5
M = 0.9


# --- утилиты ----------------------------------------------------------------
def slide_blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _run(p, text, font, size, color, bold=False, track=None):
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    if track:
        r._r.get_or_add_rPr().set("spc", str(int(track * 100)))
    return r


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def line(slide, x, y, w, color, weight=0.75):
    shp = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    shp.line.color.rgb = color
    shp.line.width = Pt(weight)
    return shp


def rect(slide, x, y, w, h, fill=None, outline=None, weight=1.0, shape=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if outline is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = outline
        shp.line.width = Pt(weight)
    return shp


def label(slide, x, y, w, text, color, size=8.5, track=2.0, align=PP_ALIGN.LEFT):
    tf = textbox(slide, x, y, w, 0.3)
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, MONO, size, color, track=track)


def page(slide, num, color):
    label(slide, 11.75, 7.0, 0.68, num, color, 9, 1.5, PP_ALIGN.RIGHT)


def stamp(slide, x, y, d, color):
    outer = rect(slide, x, y, d, d, outline=color, weight=1.0, shape=MSO_SHAPE.OVAL)
    outer.rotation = 7
    inner = rect(slide, x + 0.085, y + 0.085, d - 0.17, d - 0.17, outline=color, weight=0.5, shape=MSO_SHAPE.OVAL)
    inner.rotation = 7
    tf = textbox(slide, x, y + 0.16, d, d - 0.3, anchor=MSO_ANCHOR.MIDDLE)
    for text, font, size, bold in (("ПЕЧАТЬ", MONO, 6.0, False), ("Py", SERIF, 20, True), ("2026", MONO, 6.0, False)):
        p = tf.paragraphs[0] if text == "ПЕЧАТЬ" else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 0.92
        _run(p, text, font, size, color, bold=bold, track=1.2 if font == MONO else None)


def crop_battle():
    ASSETS.mkdir(parents=True, exist_ok=True)
    src = ROOT / "dist" / "journal-preview" / "top.png"
    dst = ASSETS / "battle.png"
    with Image.open(src) as im:
        # панель боя на правой странице лендинга, равные поля вокруг
        im.crop((796, 210, 1225, 475)).save(dst)
    return dst


# --- слайд 1: обложка ------------------------------------------------------
def slide_cover(prs):
    s = slide_blank(prs)
    bg(s, NIGHT)

    label(s, M, 0.55, 8.0, "ПОЛЕВОЙ ЖУРНАЛ", PAPER2, 9, 2.4)
    line(s, M, 0.9, 11.53, RGBColor(0x32, 0x2F, 0x28))

    tf = textbox(s, 0.82, 1.75, 9.8, 1.7)
    p = tf.paragraphs[0]
    p.line_spacing = 0.95
    _run(p, "Py", SERIF, 86, PAPER, bold=True)
    _run(p, "Dungeon", SERIF, 86, RUST_LIGHT, bold=True)

    label(s, M, 3.42, 8.0, "ПЕЧАТИ ГЛУБИНЫ", RGBColor(0x9D, 0xB3, 0xA8), 12, 4.0)

    tf = textbox(s, M, 3.98, 7.4, 1.0)
    p = tf.paragraphs[0]
    p.line_spacing = 1.28
    _run(p, "Спуск в подземелье, где каждый бой — вопрос по\nPython. Десять этажей, пять боссов, один дракон.",
         SERIF, 18, PAPER)

    tags = [("10 ЭТАЖЕЙ", 1.55), ("5 БОССОВ", 1.45), ("20–40 МИНУТ", 1.85)]
    x = M
    for text, w in tags:
        shp = rect(s, x, 5.02, w, 0.42, outline=RGBColor(0x5A, 0x51, 0x42), weight=0.75)
        tf = shp.text_frame
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, text, MONO, 8.5, PAPER2, track=1.8)
        x += w + 0.18

    stamp(s, 10.95, 1.55, 1.5, RGBColor(0xA3, 0x70, 0x5A))

    # схема спуска: только линия, деления и биомы
    x0, span = 1.15, 9.75
    step = span / 9
    line(s, x0, 6.4, span, RGBColor(0x4A, 0x45, 0x3A))
    for i in range(10):
        cx = x0 + i * step
        if i % 2 == 1:
            d = 0.15
            rect(s, cx - d / 2, 6.4 - d / 2, d, d, fill=RUST_LIGHT, shape=MSO_SHAPE.DIAMOND)
        else:
            d = 0.07
            rect(s, cx - d / 2, 6.4 - d / 2, d, d, fill=PAPER2, shape=MSO_SHAPE.OVAL)

    biomes = ["ПЕЩЕРЫ", "КАТАКОМБЫ", "ЛЁД", "ДЕМОНЫ", "ДРАКОН"]
    for j, name in enumerate(biomes):
        cx = x0 + (j * 2 + 1) * step
        tf = textbox(s, cx - 0.7, 6.55, 1.4, 0.22)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, name, MONO, 6.5, RGBColor(0x8F, 0x88, 0x7A), track=1.2)

    label(s, M, 7.02, 8.0, "PYTHON · БРАУЗЕР · БЕЗ СЕРВЕРА", RGBColor(0x6F, 0x69, 0x5C), 7.5, 2.0)
    page(s, "01", RGBColor(0x6F, 0x69, 0x5C))


# --- слайд 2: механика -----------------------------------------------------
def slide_mechanic(prs, battle: Path):
    s = slide_blank(prs)
    bg(s, PAPER)

    tf = textbox(s, M, 1.0, 6.3, 1.5)
    for i, txt in enumerate(("Вопрос — это удар.", "Ошибка — ответный.")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12
        _run(p, txt, SERIF, 33, INK, bold=True)

    steps = [
        ("01", "Вопрос", "Строки, списки, словари, функции."),
        ("02", "Задача", "Сразу после ответа — код на ту же тему."),
        ("03", "Урон", "Ошибка снимает здоровье герою, верный ответ — монстру."),
    ]
    y = 2.85
    for i, (num, title, desc) in enumerate(steps):
        label(s, M, y, 0.6, num, RUST_INK, 8.5, 2.0)
        tf = textbox(s, M + 0.55, y - 0.06, 5.6, 0.35)
        p = tf.paragraphs[0]
        _run(p, title, SERIF, 15, INK, bold=True)
        tf = textbox(s, M + 0.55, y + 0.36, 5.9, 0.7)
        p = tf.paragraphs[0]
        p.line_spacing = 1.22
        _run(p, desc, SERIF, 10.5, INK_SOFT)
        if i < len(steps) - 1:
            line(s, M, y + 1.22, 6.0, PAPER2, 0.6)
        y += 1.38

    # снимок без подписи
    px, py, pw = 7.03, 1.3, 5.4
    with Image.open(battle) as im:
        ph = pw * im.size[1] / im.size[0]
    rect(s, px - 0.09, py - 0.09, pw + 0.18, ph + 0.18, fill=NIGHT2)
    s.shapes.add_picture(str(battle), Inches(px), Inches(py), width=Inches(pw))
    rect(s, px, py, pw, ph, outline=NIGHT2, weight=1.25)

    tape = rect(s, px + 0.18, py - 0.16, 1.05, 0.24, fill=PAPER3)
    tape.rotation = -3.5

    tf = textbox(s, px, py + ph + 0.28, pw, 0.4)
    p = tf.paragraphs[0]
    _run(p, "Проверка кода идёт в браузере.", SERIF, 11.5, INK_SOFT)

    page(s, "02", INK_SOFT)


# --- слайд 3: опись -------------------------------------------------------
def slide_ledger(prs):
    s = slide_blank(prs)
    bg(s, PAPER)

    tf = textbox(s, M, 1.05, 11.4, 0.9)
    p = tf.paragraphs[0]
    _run(p, "Десять этажей. Пять боссов. Один дракон.", SERIF, 30, INK, bold=True)

    tf = textbox(s, M, 1.74, 11.0, 0.4)
    p = tf.paragraphs[0]
    _run(p, "Кампания «Печати глубины»: пять глав, деревня, журнал и финал.", SERIF, 12.5, INK_SOFT)

    rows = [
        ("ЭТАЖЕЙ", "10"),
        ("БИОМОВ", "5"),
        ("БОССОВ", "5"),
        ("МИНУТ НА ПОХОД", "20–40"),
        ("СЕРВЕР", "не нужен"),
    ]
    y = 2.55
    for name, value in rows:
        label(s, M, y + 0.16, 4.0, name, INK_SOFT, 9, 2.2)
        tf = textbox(s, 9.4, y, 3.03, 0.6)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        size = 26 if value[0].isdigit() else 17
        _run(p, value, MONO, size, INK, track=0.6 if size == 26 else 0)
        line(s, M, y + 0.66, 11.53, PAPER2, 0.6)
        y += 0.7

    tf = textbox(s, M, 6.28, 11.5, 0.5)
    p = tf.paragraphs[0]
    _run(p, "Python через Pyodide, пиксель-арт 32×32 (Aseprite, открытые тайлы DCSS), сохранения — в localStorage.",
         SERIF, 11.5, INK_SOFT)

    page(s, "03", INK_SOFT)


def main():
    battle = crop_battle()
    prs = Presentation()
    prs.slide_width = Emu(int(W * 914400))
    prs.slide_height = Emu(int(H * 914400))

    slide_cover(prs)
    slide_mechanic(prs, battle)
    slide_ledger(prs)

    for slide, idx in zip(prs.slides, (1, 2, 3)):
        tf = slide.notes_slide.notes_text_frame
        tf.text = NOTES[idx]

    prs.save(OUT)
    print(f"saved: {OUT}")


if __name__ == "__main__":
    main()
